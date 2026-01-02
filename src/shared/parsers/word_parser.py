"""Word document parser implementation."""

import asyncio
import logging
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import List

from .file_parser import (
    FileParser, ParseResult, PrivacyTextSegment, FileParserError
)
from .parser_config import ParserConfig
from ..constant.enums import ParseMode, FileType

logger = logging.getLogger(__name__)


class WordParser(FileParser):
    """Parser for Word documents (.docx and .doc) optimized for privacy detection."""

    SUPPORTED_EXTENSIONS = ['docx', 'doc']

    # Confidence scores
    CONFIDENCE_DOCX_TEXT = 0.9
    CONFIDENCE_DOCX_TABLE = 0.95
    CONFIDENCE_DOC_ANTIWORD = 0.8
    CONFIDENCE_DOC_FALLBACK = 0.7

    def __init__(self, config: ParserConfig):
        super().__init__(config)

    @property
    def supported_file_types(self) -> List[FileType]:
        return [FileType.WORD_DOCX, FileType.WORD_DOC]

    @property
    def supported_modes(self) -> List[ParseMode]:
        return [ParseMode.LOCAL]  # Only local parsing supported

    async def parse_local(self, file_path: Path, file_type: FileType) -> ParseResult:
        """Parse Word document using local processing."""
        try:
            # Run parsing in thread pool to avoid blocking
            loop = asyncio.get_running_loop()
            with ThreadPoolExecutor() as executor:
                segments = await loop.run_in_executor(
                    executor, self._extract_privacy_segments, file_path, file_type
                )

            # Filter for privacy-relevant segments
            privacy_segments = self.filter_privacy_segments(segments)

            metadata = {
                "total_segments": len(segments),
                "privacy_segments": len(privacy_segments),
                "parser_used": "python-docx" if file_type == FileType.WORD_DOCX else "antiword/fallback",
                "file_info": self.get_file_info(file_path)
            }

            return ParseResult(
                segments=privacy_segments,
                metadata=metadata,
                success=True
            )

        except Exception as e:
            logger.error(f"Failed to parse Word document {file_path}: {e}", exc_info=True)
            raise FileParserError(f"Word document parsing failed: {str(e)}")

    async def parse_cloud(self, file_path: Path, file_type: FileType) -> ParseResult:
        """Parse Word document using cloud services (placeholder)."""
        # Placeholder for cloud document parsing (e.g., Google Document AI)
        raise FileParserError("Cloud Word parsing not implemented yet")

    def _extract_privacy_segments(self, file_path: Path, file_type: FileType) -> List[PrivacyTextSegment]:
        """Extract privacy-relevant text segments from Word document."""
        if file_type == FileType.WORD_DOCX:
            return self._extract_from_docx(file_path)
        elif file_type == FileType.WORD_DOC:
            return self._extract_from_doc(file_path)
        else:
            raise FileParserError(f"Unsupported Word file type: {file_type}")

    def _extract_from_docx(self, file_path: Path) -> List[PrivacyTextSegment]:
        """Extract text segments from .docx files."""
        segments = []

        try:
            from docx import Document
        except ImportError:
            logger.error("python-docx not available. Install with: pip install python-docx")
            return segments

        try:
            doc = Document(str(file_path))

            # Extract text from paragraphs
            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if text:
                    segments.append(PrivacyTextSegment(
                        text=text,
                        confidence=self.CONFIDENCE_DOCX_TEXT,
                        location={"paragraph": i, "source": "docx_paragraph"},
                        source_type="word_doc",
                        privacy_indicators=[]
                    ))

            # Extract text from tables
            for table_idx, table in enumerate(doc.tables):
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        text = cell.text.strip()
                        if text:
                            segments.append(PrivacyTextSegment(
                                text=text,
                                confidence=self.CONFIDENCE_DOCX_TABLE,
                                location={
                                    "table": table_idx,
                                    "row": row_idx,
                                    "col": col_idx,
                                    "source": "docx_table_cell"
                                },
                                source_type="word_doc",
                                privacy_indicators=[]
                            ))

        except Exception as e:
            logger.error(f"DOCX parsing failed for {file_path}: {e}")

        return segments

    def _extract_from_doc(self, file_path: Path) -> List[PrivacyTextSegment]:
        """Extract text segments from .doc files."""
        segments = self._try_antiword(file_path)

        if not segments:
            logger.debug(f"Antiword failed or returned empty for {file_path}, trying fallback.")
            segments = self._try_docx_fallback(file_path)

        return segments

    def _try_antiword(self, file_path: Path) -> List[PrivacyTextSegment]:
        """Try parsing .doc with antiword utility."""
        segments = []
        try:
            result = subprocess.run(
                ['antiword', str(file_path)],
                capture_output=True,
                text=True,
                timeout=30,
                check=False
            )

            if result.returncode == 0 and result.stdout:
                paragraphs = result.stdout.split('\n\n')
                for i, para in enumerate(paragraphs):
                    text = para.strip()
                    if text:
                        segments.append(PrivacyTextSegment(
                            text=text,
                            confidence=self.CONFIDENCE_DOC_ANTIWORD,
                            location={"paragraph": i, "source": "doc_antiword"},
                            source_type="word_doc",
                            privacy_indicators=[]
                        ))
            else:
                if result.stderr:
                    logger.debug(f"Antiword stderr: {result.stderr}")

        except FileNotFoundError:
            logger.warning("antiword utility not found in system path.")
        except Exception as e:
            logger.warning(f"Antiword parsing error: {e}")

        return segments

    def _try_docx_fallback(self, file_path: Path) -> List[PrivacyTextSegment]:
        """Fallback: try treating .doc as .docx (sometimes works for renamed files)."""
        segments = []
        try:
            from docx import Document
            doc = Document(str(file_path))

            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if text:
                    segments.append(PrivacyTextSegment(
                        text=text,
                        confidence=self.CONFIDENCE_DOC_FALLBACK,
                        location={"paragraph": i, "source": "doc_fallback"},
                        source_type="word_doc",
                        privacy_indicators=[]
                    ))
        except Exception:
            # Expected failure for actual binary .doc files
            pass

        return segments
