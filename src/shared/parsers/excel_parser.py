"""Excel file parser implementation."""

import asyncio
import logging
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import List

from .file_parser import (
    FileParser, ParseResult, PrivacyTextSegment, FileParserError
)
from .parser_config import ParserConfig
from ..constant.enums import FileType, ParseMode


class ExcelParser(FileParser):
    """Parser for Excel files (.xlsx and .xls) optimized for privacy detection."""

    SUPPORTED_EXTENSIONS = ['xlsx', 'xls']

    def __init__(self, config: ParserConfig):
        super().__init__(config)
        self.logger = logging.getLogger(self.__class__.__name__)

    @property
    def supported_file_types(self) -> List[FileType]:
        return [FileType.EXCEL_XLSX, FileType.EXCEL_XLS]

    @property
    def supported_modes(self) -> List[ParseMode]:
        return [ParseMode.LOCAL]  # Only local parsing supported

    async def parse_local(self, file_path: Path, file_type: FileType) -> ParseResult:
        """Parse Excel file using local processing."""
        try:
            # Run parsing in thread pool to avoid blocking
            loop = asyncio.get_event_loop()
            with ThreadPoolExecutor() as executor:
                segments = await loop.run_in_executor(
                    executor, self._extract_privacy_segments, file_path, file_type
                )

            # Filter for privacy-relevant segments
            privacy_segments = self.filter_privacy_segments(segments)

            metadata = {
                "total_segments": len(segments),
                "privacy_segments": len(privacy_segments),
                "parser_used": "openpyxl" if file_type == FileType.EXCEL_XLSX else "xlrd",
                "file_info": self.get_file_info(file_path)
            }

            return ParseResult(
                segments=privacy_segments,
                metadata=metadata,
                success=True
            )

        except Exception as e:
            self.logger.error(f"Failed to parse Excel file {file_path}: {e}")
            raise FileParserError(f"Excel parsing failed: {str(e)}")

    async def parse_cloud(self, file_path: Path, file_type: FileType) -> ParseResult:
        """Parse Excel file using cloud services (placeholder)."""
        # Placeholder for cloud Excel parsing (e.g., Google Sheets API)
        raise FileParserError("Cloud Excel parsing not implemented yet")

    def _extract_privacy_segments(self, file_path: Path, file_type: FileType) -> List[PrivacyTextSegment]:
        """Extract privacy-relevant text segments from Excel file."""
        segments = []

        try:
            if file_type == FileType.EXCEL_XLSX:
                segments = self._extract_from_xlsx(file_path)
            elif file_type == FileType.EXCEL_XLS:
                segments = self._extract_from_xls(file_path)

        except Exception as e:
            self.logger.error(f"Excel parsing failed: {str(e)}")

        return segments

    def _extract_from_xlsx(self, file_path: Path) -> List[PrivacyTextSegment]:
        """Extract text segments from .xlsx file."""
        segments = []

        try:
            import openpyxl
        except ImportError:
            self.logger.error("openpyxl not available for .xlsx parsing")
            return segments

        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)

            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]

                for row_idx, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                    for col_idx, cell_value in enumerate(row, 1):
                        if cell_value is not None:
                            cell_text = str(cell_value).strip()
                            if cell_text:
                                segment = PrivacyTextSegment(
                                    text=cell_text,
                                    confidence=0.95,  # Excel cells often contain structured data
                                    location={
                                        "sheet": sheet_name,
                                        "row": row_idx,
                                        "col": col_idx,
                                        "source": "xlsx_cell"
                                    },
                                    source_type="excel",
                                    privacy_indicators=[]
                                )
                                segments.append(segment)

        except Exception as e:
            self.logger.error(f"XLSX parsing failed: {str(e)}")

        return segments

    def _extract_from_xls(self, file_path: Path) -> List[PrivacyTextSegment]:
        """Extract text segments from .xls file."""
        segments = []

        try:
            import xlrd
        except ImportError:
            self.logger.error("xlrd not available for .xls parsing")
            return segments

        try:
            workbook = xlrd.open_workbook(str(file_path))

            for sheet_idx in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(sheet_idx)
                sheet_name = sheet.name

                for row_idx in range(sheet.nrows):
                    for col_idx in range(sheet.ncols):
                        cell = sheet.cell(row_idx, col_idx)
                        if cell.ctype != xlrd.XL_CELL_EMPTY:
                            cell_value = cell.value
                            if cell_value is not None:
                                cell_text = str(cell_value).strip()
                                if cell_text:
                                    segment = PrivacyTextSegment(
                                        text=cell_text,
                                        confidence=0.9,  # Slightly lower confidence for older format
                                        location={
                                            "sheet": sheet_name,
                                            "row": row_idx + 1,
                                            "col": col_idx + 1,
                                            "source": "xls_cell"
                                        },
                                        source_type="excel",
                                        privacy_indicators=[]
                                    )
                                    segments.append(segment)

        except Exception as e:
            self.logger.error(f"XLS parsing failed: {str(e)}")

        return segments

